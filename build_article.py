# -*- coding: utf-8 -*-
"""
Build the Economic Archive (Narodnostopanski Arhiv) article (.docx) following the
journal template. Topic: digital transformation in SMEs (dominant theme of the
supplied collection), management-oriented, adapted to emerging Asia. Revised
twice for peer-review readiness (estimable specification, measurable capability,
honest novelty framing, narrative review, country vignette, recent references,
language).
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FONT = "Times New Roman"
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for s in doc.sections:
    s.top_margin = Inches(1); s.bottom_margin = Inches(1)
    s.left_margin = Inches(1.18); s.right_margin = Inches(1.18)


def set_run(r, size=14, bold=False, italic=False):
    r.font.name = FONT; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def para(text="", *, size=14, bold=False, italic=False, align="justify",
         spacing=1.5, space_after=6, first_indent=0.0):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE if isinstance(spacing, float) else WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(space_after); pf.space_before = Pt(0)
    if first_indent:
        pf.first_line_indent = Inches(first_indent)
    if text:
        set_run(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def heading(t, *, size=14, sub=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8 if sub else 10); p.paragraph_format.space_after = Pt(4 if sub else 6)
    set_run(p.add_run(t), size=size, bold=True, italic=sub)
    return p


def equation(txt, number):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(txt), italic=True)
    set_run(p.add_run("\t\t\t\t\t(%d)" % number))
    return p


def figure(path, caption, source, width=6.0):
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(6)
    pic.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = 1.0; cap.paragraph_format.space_after = Pt(2)
    set_run(cap.add_run(caption), size=12, bold=True)
    src = doc.add_paragraph(); src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    src.paragraph_format.line_spacing = 1.0; src.paragraph_format.space_after = Pt(10)
    set_run(src.add_run(source), size=11, italic=True)


def table(title_no, title, headers, rows, source):
    para(title_no, bold=True, align="left", spacing=1.0, space_after=2)
    para(title, bold=True, size=12, align="left", spacing=1.0, space_after=4)
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, headers):
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(c.paragraphs[0].add_run(h), size=12, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for cell, txt in zip(cells, row):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run(cell.paragraphs[0].add_run(txt), size=12)
    para(source, italic=True, size=11, align="left", spacing=1.0, space_after=10)


# ===================== FRONT MATTER =====================
para("Digital Transformation and the Performance of Small and Medium-Sized "
     "Enterprises: A Management Framework and Evidence from Emerging Asia".upper(),
     bold=True, align="center", spacing=1.0, space_after=10)
para("[Author Name], [e-mail]", bold=True, align="center", spacing=1.0, space_after=2)
para("[Department]", bold=True, align="center", spacing=1.0, space_after=2)
para("[Institution]", bold=True, align="center", spacing=1.0, space_after=12)

abstract = (
    "Abstract: Digital transformation is the central management challenge for small and "
    "medium-sized enterprises (SMEs), which dominate the economies of emerging Asia yet adopt "
    "advanced digital technologies more slowly than large firms. This study develops a management "
    "framework explaining how SMEs convert digital technologies into superior performance. Building "
    "on the dynamic-capabilities perspective and on the distinction between digitisation, "
    "digitalisation and digital transformation, the framework links external drivers to the "
    "firm\u2019s sensing, seizing and reconfiguring capabilities, to the depth of the transformation "
    "undertaken, and to performance, with a reinvestment feedback loop. The framework\u2019s "
    "contribution is integrative rather than metric: it embeds firm-level digital-adoption measures "
    "in a capability-conditioned model and specifies a moderated relationship \u2014 a "
    "digital-transformation intensity (DTI) measure interacting with dynamic capability \u2014 that "
    "can be empirically tested once both constructs are operationalised, which the study sets out "
    "using established adoption items and a validated dynamic-capabilities scale. The framework is "
    "built through a narrative review of the literature, illustrated with a numerical example of the "
    "moderation logic, and assessed in light of secondary evidence from emerging Asia, including the "
    "ASEAN region and a country vignette of Uzbekistan. The evidence suggests that the binding "
    "constraint on SME performance is managerial and human capability rather than connectivity, and "
    "the analysis yields managerial and policy recommendations on capability building, finance, "
    "digital skills and shared digital infrastructure."
)
para(abstract, spacing=1.0, space_after=8)
para("Key words: digital transformation; small and medium-sized enterprises; dynamic capabilities; "
     "firm performance; emerging Asia.", spacing=1.0, space_after=4)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(12)
set_run(p.add_run("JEL: M15, M21, O33, L25, O18, O53."), bold=True)

# ===================== INTRODUCTION =====================
heading("Introduction")
para(
    "Digital technologies \u2014 cloud computing, data analytics, artificial intelligence, "
    "e-commerce platforms and connected devices \u2014 have moved from being a support function to "
    "being the principal arena of competition. For enterprises of every size, the management "
    "question is no longer whether to adopt such technologies but how to convert them into durable "
    "improvements in performance. This question is most acute for small and medium-sized enterprises "
    "(SMEs). In emerging Asia they are the backbone of the economy: in the member states of the "
    "Association of Southeast Asian Nations (ASEAN), micro, small and medium-sized enterprises make "
    "up on average more than ninety-seven per cent of all firms and provide the large majority of "
    "employment (Asian Development Bank, 2023). Yet they adopt advanced digital technologies more "
    "slowly than large enterprises and face tighter constraints on finance, skills and managerial "
    "capacity (OECD, 2021; Yoshino & Taghizadeh-Hesary, 2016).",
    first_indent=0.30)
para(
    "The literature has clarified what digital transformation is, but it has paid less attention to "
    "the managerial mechanisms through which a resource-constrained SME in an emerging economy turns "
    "technology into results. Digital transformation is an organisational change process in which "
    "digital technologies reshape value creation, operations and the business model itself (Vial, "
    "2019; Verhoef et al., 2021), and recent comprehensive reviews stress that its outcomes hinge on "
    "managerial and organisational factors rather than on technology adoption alone (Hanelt et al., "
    "2021; Kraus et al., 2022). Whether the process raises performance depends on the capabilities "
    "the firm can mobilise around it (Teece, 2007; Warner & W\u00e4ger, 2019), and SME-focused "
    "reviews report that transformation levels in smaller firms remain low and that the field still "
    "lacks process-level guidance (de Mattos et al., 2024; Sagala & \u0150ri, 2024).",
    first_indent=0.30)
para(
    "This study addresses that gap with a management framework connecting the external drivers of "
    "digital transformation, the dynamic capabilities of the firm, the depth of the transformation "
    "undertaken and the resulting performance, closed by a reinvestment feedback loop. Its "
    "contribution is integrative rather than the proposal of a new metric. Firm-level digital "
    "adoption is already measured \u2014 for example by the World Bank Enterprise Surveys, by the "
    "European digital-intensity indicators and in survey studies such as Teng, Wu and Yang (2022). "
    "What the framework adds is, first, to embed such adoption measures in an explicit "
    "capability-conditioned causal structure; second, to state the often-asserted claim that "
    "\u201ctechnology alone is not enough\u201d as a moderated relationship with a clear empirical "
    "test (the interaction between digital-transformation intensity and dynamic capability); and "
    "third, to apply the framework to emerging Asia \u2014 the ASEAN region and the under-researched "
    "economies of Central Asia, with Uzbekistan as a country vignette \u2014 where firm-level "
    "evidence is comparatively thin.",
    first_indent=0.30)
para(
    "The remainder of the article is organised as follows. Section 1 reviews the literature and "
    "defines the core constructs. Section 2 presents the framework, the measures and the estimable "
    "specification. Section 3 describes the data, the operationalisation of the constructs and the "
    "estimation strategy. Section 4 reports a numerical illustration of the moderation logic and the "
    "regional evidence, including the Uzbekistan vignette. Section 5 discusses the management and "
    "policy implications. The final section concludes.",
    first_indent=0.30)

# ===================== 1. LITERATURE =====================
heading("1. Theoretical background and literature review")
para(
    "Research distinguishes three increasingly profound stages of digital change. Digitisation is "
    "the conversion of analogue information into digital form; digitalisation is the use of digital "
    "technologies to improve existing processes; and digital transformation is the deeper, "
    "firm-wide change in which technology reconfigures value creation and the business model "
    "(Verhoef et al., 2021). Vial (2019) defines digital transformation as a process in which "
    "digital technologies trigger strategic responses that alter value-creation paths while the firm "
    "manages structural change and organisational barriers; large-scale reviews confirm this "
    "organisational-change reading and the centrality of strategy and leadership (Hanelt et al., "
    "2021; Kraus et al., 2022). The implication for management is that technology alone does not "
    "deliver value; the value comes from the organisational changes that accompany it.",
    first_indent=0.30)
para(
    "The dynamic-capabilities perspective explains why firms facing the same technologies achieve "
    "different results. Teece (2007) decomposes the capacity to sustain performance into sensing "
    "opportunities and threats, seizing them through investment and business-model choices, and "
    "reconfiguring (transforming) the resource base accordingly. Warner and W\u00e4ger (2019) show "
    "that digital transformation is precisely an ongoing process of strategic renewal in which these "
    "capabilities are built and rebuilt, and Kump, Engelmann, Kessler and Schweiger (2019) provide a "
    "validated survey scale that measures sensing, seizing and transforming capacities separately. "
    "Bharadwaj, El Sawy, Pavlou and Venkatraman (2013) argue that digital strategy must be fused "
    "with business strategy; Westerman, Bonnet and McAfee (2014) add that the binding constraint is "
    "usually leadership and management capability rather than technology as such; and Fitzgerald, "
    "Kruschwitz, Bonnet and Welch (2014) document that firms recognise digital technology as a "
    "strategic imperative while struggling to execute it.",
    first_indent=0.30)
para(
    "A further strand explains how the adoption of a specific technology is decided, and it feeds "
    "directly into the framework\u2019s measures. The technology acceptance model relates adoption "
    "to perceived usefulness and perceived ease of use (Davis, 1989), and the unified theory of "
    "acceptance and use of technology adds performance and effort expectancy, social influence and "
    "facilitating conditions (Venkatesh, Morris, Davis & Davis, 2003); both inform the items that "
    "make up the adoption levels aggregated in the digital-transformation intensity measure below, "
    "and managerial acceptance is itself part of the sensing and seizing capability. At the "
    "organisational level, the technology\u2013organisation\u2013environment framework situates "
    "adoption in the interplay of technological readiness, organisational resources and the external "
    "environment (Tornatzky & Fleischer, 1990), which maps onto the framework\u2019s capability and "
    "driver blocks respectively. For SMEs in emerging economies these models matter because the "
    "adoption decision is concentrated in a small management team and is highly sensitive to skills, "
    "cost and the surrounding digital ecosystem.",
    first_indent=0.30)
para(
    "Recent empirical and policy work sharpens the picture. Teng et al. (2022), using a "
    "structural-equation model on 335 SMEs, find that digital technology, employee digital skills "
    "and a digital-transformation strategy are jointly and positively associated with transformation "
    "and, through it, with performance \u2014 evidence consistent with a capability-conditioned view "
    "and with an interaction between technology and capability. The policy-oriented literature "
    "documents the SME-specific barriers \u2014 scarce skills, limited finance and uncertainty about "
    "returns (OECD, 2021; Yoshino & Taghizadeh-Hesary, 2016) \u2014 and Nambisan, Wright and Feldman "
    "(2019) note that digital technologies also reshape innovation and entrepreneurship themselves, "
    "raising the premium on digital capability. This body of work motivates a framework in which "
    "drivers, firm-level capabilities and the depth of transformation jointly determine performance, "
    "and a specification that can be tested.",
    first_indent=0.30)

# ===================== 2. FRAMEWORK =====================
heading("2. Framework, measures and estimable specification")
para(
    "The framework links four blocks in sequence, with a feedback loop (Figure 1). The first block "
    "is the set of external drivers: market and competitive pressure, the availability and cost of "
    "digital technologies, and the policy environment, including funding and regulation. Drivers "
    "create the incentive and the opportunity to transform, but they do not determine the outcome.",
    first_indent=0.30)
para(
    "The second block is the firm\u2019s dynamic capabilities. Following Teece (2007), an SME senses "
    "digital opportunities and threats, seizes them by investing in technology and adjusting its "
    "business model, and reconfigures its processes, skills and structure so that the technology is "
    "actually used. The third block is the depth of the transformation undertaken, from digitisation "
    "through digitalisation to business-model change (Verhoef et al., 2021). The fourth block is "
    "performance: productivity, growth and export reach, and resilience. A feedback loop closes the "
    "framework: performance gains generate resources that are reinvested in further capability "
    "building and technology.",
    first_indent=0.30)

figure("figure1_framework.png",
       "Figure 1. A management framework for digital transformation in SMEs",
       "Source: Authors\u2019 elaboration based on Teece (2007), Vial (2019) and Verhoef et al. (2021).")

heading("2.1. Measuring digital-transformation intensity and capability", sub=True)
para(
    "Two constructs must be measured. Digital-transformation intensity (DTI) summarises the digital "
    "technologies and practices a firm actually uses, as a weighted composite",
    first_indent=0.30)
equation("DTI = \u03a3\u1d62 w\u1d62 \u00b7 a\u1d62 ,   \u03a3\u1d62 w\u1d62 = 1 ,   a\u1d62 \u2208 [0, 1] ,", 1)
para(
    "where a\u1d62 is the adoption level of digital element i \u2014 connectivity, a website or "
    "e-commerce channel, digital payments, cloud services, enterprise software, data analytics, "
    "artificial intelligence and digital-security measures \u2014 and w\u1d62 is its weight. The "
    "elements and items follow established firm-level instruments (the World Bank Enterprise Surveys "
    "and European digital-intensity indicators) and the acceptance literature (Davis, 1989; "
    "Venkatesh et al., 2003); the measure is therefore not a new statistic but a transparent "
    "aggregation of existing ones for use inside the framework.",
    first_indent=0.30)
para(
    "Capability (C) is operationalised, rather than left undefined, as the firm\u2019s dynamic "
    "capability measured on the validated scale of Kump et al. (2019), which captures sensing (S\u2081), "
    "seizing (S\u2082) and transforming (S\u2083) capacities through multi-item managerial-survey "
    "constructs. A composite capability index can be formed as",
    first_indent=0.30)
equation("C = (S\u2081 + S\u2082 + S\u2083) / 3 ,   C \u2208 [0, 1] ,", 2)
para(
    "after rescaling each sub-scale to the unit interval, with observable complements \u2014 "
    "managerial digital skills, the presence of an explicit digital strategy, prior technology "
    "projects and staff training \u2014 available as proxies where survey access is limited. Both "
    "DTI and C are thus measurable with existing, peer-validated instruments.",
    first_indent=0.30)

heading("2.2. The estimable specification", sub=True)
para(
    "Let \u0394P denote the change in a performance measure (for example labour productivity, sales "
    "growth or export intensity). The framework\u2019s central claim \u2014 that adoption pays off "
    "only when matched by capability \u2014 is a moderation hypothesis, which is estimated in the "
    "standard interaction form rather than asserted as a pure product:",
    first_indent=0.30)
equation("\u0394P = \u03b2\u2080 + \u03b2\u2081\u00b7DTI + \u03b2\u2082\u00b7C + \u03b2\u2083\u00b7(DTI\u00d7C) + \u03b3\u00b7X + \u03b5 ,", 3)
para(
    "where X is a vector of controls (firm size, age, sector, region) and \u03b5 is the error term. "
    "The interaction term DTI\u00d7C is the object of interest: a positive \u03b2\u2083 means the "
    "marginal performance return to digital adoption, \u2202\u0394P/\u2202DTI = \u03b2\u2081 + "
    "\u03b2\u2083\u00b7C, rises with capability \u2014 the complementarity between technology and "
    "capability familiar from absorptive-capacity and complementarity theory and consistent with the "
    "SEM evidence of Teng et al. (2022). The additive terms are retained precisely because theory "
    "does not require a pure product; equation (3) nests the simpler case. When capability is the "
    "overwhelming constraint (\u03b2\u2081 \u2248 0, \u03b2\u2082 \u2248 0), equation (3) collapses to "
    "the strong-complementarity special case",
    first_indent=0.30)
equation("\u0394P \u2248 \u03b2\u2083\u00b7(DTI\u00d7C) ,", 4)
para(
    "which is the form used only for the didactic illustration in Section 4.1. The coefficients "
    "\u03b2\u2080\u2013\u03b2\u2083 are parameters to be estimated, not assumed; the illustration "
    "fixes them only to visualise the qualitative shape of the moderation, and makes no claim about "
    "their magnitude.",
    first_indent=0.30)

# ===================== 3. DATA & METHODS =====================
heading("3. Data, operationalisation and estimation strategy")
para(
    "This article is conceptual with an empirical research design specified for testing; it does not "
    "itself collect primary data. The framework was developed through a narrative (non-exhaustive) "
    "review of the literature on digital transformation, dynamic capabilities and technology "
    "adoption, drawing on highly cited reviews and SME-specific studies; it does not claim the "
    "exhaustive search protocol, screening counts or PRISMA reporting of a systematic review, and is "
    "presented as a narrative synthesis rather than a systematic one.",
    first_indent=0.30)
para(
    "For testing the specification in equation (3), the constructs are operationalised as follows. "
    "DTI is computed from firm-level adoption items of the kind collected in the World Bank "
    "Enterprise Surveys (for example use of a website or e-commerce, e-mail with clients and "
    "suppliers, and \u2014 in recent rounds \u2014 cloud and digital payments) and the European "
    "digital-intensity indicators, with equal weights as a baseline and entropy- or expert-derived "
    "weights as robustness checks. Capability C is measured with the Kump et al. (2019) "
    "sensing\u2013seizing\u2013transforming scale administered to owner-managers, with the observable "
    "proxies noted in Section 2.1 used where a full survey is infeasible. Performance \u0394P is "
    "taken from accounts or survey self-reports (productivity, sales growth, export intensity). "
    "Equation (3) is then estimated by ordinary least squares with robust standard errors, or by "
    "structural-equation modelling when latent constructs are modelled directly, and the moderation "
    "is assessed from the sign, size and significance of \u03b2\u2083 and from marginal-effect plots "
    "of \u2202\u0394P/\u2202DTI across the range of C.",
    first_indent=0.30)
para(
    "Pending such firm-level estimation, the framework is assessed in light of secondary evidence "
    "for emerging Asia drawn from official and institutional sources: the Asian Development Bank "
    "(2023, 2024) and the OECD and ERIA (2024) for the ASEAN region; Yoshino and Taghizadeh-Hesary "
    "(2016) for the structural constraints on Asian SMEs; and, for the Uzbekistan vignette, the "
    "World Bank (2023, 2025) and a UNESCAP (2025) foresight initiative. Quantifying "
    "\u03b2\u2080\u2013\u03b2\u2083 on firm-level data, including a primary survey of Uzbek and ASEAN "
    "SMEs, is the principal direction for further work.",
    first_indent=0.30)

# ===================== 4. RESULTS =====================
heading("4. Illustration and regional evidence")

heading("4.1. Numerical illustration of the moderation logic", sub=True)
para(
    "Before turning to the evidence, a numerical example clarifies what the moderation in equation "
    "(3) implies; it is a didactic device, not new empirical information, and simply traces the "
    "strong-complementarity form of equation (4) for transparent inputs. Three representative SME "
    "profiles are used: a micro service firm with full connectivity but no analytics or artificial "
    "intelligence (DTI \u2248 0.33), a small manufacturer with moderate adoption (DTI \u2248 0.54) "
    "and a digitally advanced medium-sized exporter (DTI \u2248 0.80). Holding the slope at a "
    "common value, the implied performance gain is plotted for a high capability level (C = 0.8) and "
    "a low one (C = 0.3) in Figure 2, with the underlying figures in Table 1.",
    first_indent=0.30)

figure("figure2_dti_application.png",
       "Figure 2. Numerical illustration: performance gain by DTI level under high and low capability",
       "Source: Authors\u2019 illustration of equation (4) for didactic purposes; the values are inputs, "
       "not estimates or survey data.",
       width=5.6)

table("Table 1",
      "Numerical illustration of the DTI\u00d7capability moderation (didactic, not empirical)",
      ["SME profile", "DTI", "Gain at C = 0.8", "Gain at C = 0.3"],
      [["Micro service firm", "0.33", "0.26", "0.10"],
       ["Small manufacturer", "0.54", "0.43", "0.16"],
       ["Medium exporter", "0.80", "0.64", "0.24"]],
      "Source: Authors\u2019 illustration of equation (4) with a common slope set to one. The values "
      "are didactic inputs, not estimates or survey data.")

para(
    "The example conveys one point only: under the moderated specification, the performance return "
    "to a given level of digital adoption is far larger when capability is high than when it is low, "
    "so a digitally advanced firm with weak capability gains little relative to its potential. This "
    "is the qualitative content of the claim that capability, not technology, is the binding "
    "constraint; establishing its magnitude requires estimating equation (3) on firm-level data, as "
    "set out in Section 3.",
    first_indent=0.30)

heading("4.2. Regional evidence from emerging Asia", sub=True)
para(
    "Across the ASEAN member states, micro, small and medium-sized enterprises account for the "
    "overwhelming majority of firms and the bulk of employment, and the regional digital economy is "
    "expanding rapidly (Asian Development Bank, 2023). The constraints on their digital "
    "transformation are, however, systematic. Yoshino and Taghizadeh-Hesary (2016) show that Asian "
    "SMEs are held back by limited access to finance, the absence of comprehensive databases, low "
    "research and development spending and underdeveloped sales channels \u2014 conditions that "
    "depress the capability term in the framework. The OECD and ERIA (2024) confirm that, although "
    "ASEAN governments have strengthened SME policy frameworks, the digitalisation of smaller firms "
    "remains uneven and skills and finance are recurring bottlenecks. The regional evidence "
    "therefore suggests that connectivity and policy ambition are necessary but not sufficient: the "
    "binding constraint is capability.",
    first_indent=0.30)

heading("4.3. Country vignette: Uzbekistan", sub=True)
para(
    "Uzbekistan illustrates both the opportunity and the constraint. SMEs are central to the "
    "economy: the World Bank (2025) reports that micro, small and medium-sized enterprises account "
    "for over ninety per cent of businesses, about seventy-five per cent of employment and roughly "
    "fifty-five per cent of gross domestic product, and the national statistics authorities record "
    "more than 1.2 million small businesses in early 2025. Connectivity and e-commerce have grown "
    "quickly from a low base; the World Bank (2023) reports that the e-commerce market expanded "
    "roughly fivefold between 2018 and 2022, exceeding half a billion United States dollars by 2023, "
    "and national policy under the Digital Uzbekistan 2030 strategy aims to turn the country into a "
    "regional information-technology hub with higher technology exports and large-scale job creation.",
    first_indent=0.30)
para(
    "The deeper transformation of enterprises nonetheless lags the roll-out of infrastructure. A "
    "UNESCAP (2025) foresight initiative reports that only about ten per cent of SMEs were registered "
    "on the national digital public-services portal, suggesting that the uptake of even basic "
    "digital government services \u2014 let alone analytics or artificial intelligence \u2014 remains "
    "shallow. Read through the framework, Uzbek SMEs increasingly possess the drivers and the "
    "connectivity but not yet the capability \u2014 skills, finance and managerial capacity \u2014 "
    "that converts adoption into performance. Table 2 summarises the regional and Uzbekistan "
    "evidence.",
    first_indent=0.30)

table("Table 2",
      "Digital transformation of SMEs in emerging Asia: selected secondary evidence",
      ["Indicator / observation", "Evidence", "Source"],
      [["MSME share of enterprises, ASEAN", "On average > 97% of firms", "ADB (2023)"],
       ["Main SME constraints in Asia", "Finance; databases; R&D; sales channels", "Yoshino & Taghizadeh-Hesary (2016)"],
       ["Uzbekistan \u2013 MSME role", "> 90% of firms; ~75% jobs; ~55% of GDP", "World Bank (2025)"],
       ["Uzbekistan \u2013 e-commerce", "~5\u00d7 growth 2018\u20132022; > USD 0.5 bn (2023)", "World Bank (2023)"],
       ["Uzbekistan \u2013 SMEs on e-gov portal", "Reportedly ~10% registered", "UNESCAP (2025)"],
       ["Binding constraint (framework)", "Capability: skills, finance, management", "This study"]],
      "Source: Compiled by the authors from the cited institutional sources.")

# ===================== 5. IMPLICATIONS =====================
heading("5. Management and policy implications")
para(
    "For the management of the individual SME, the framework implies that investment should be "
    "sequenced according to capability rather than technology fashion. Because the performance "
    "return to adoption rises with capability (equation 3), a firm with weak capabilities gains more "
    "from building sensing and absorptive capacity \u2014 digital skills, data literacy and a clear "
    "digital element of strategy \u2014 than from acquiring advanced tools it cannot embed. Adoption "
    "should follow the sequence of the framework: secure connectivity and a market-facing digital "
    "channel, including the digital-payment and e-commerce platforms spreading fastest in Asian "
    "markets; move core processes to cloud-based and enterprise software; and only then layer on "
    "analytics and artificial intelligence, with cybersecurity and data governance treated as a "
    "precondition. Digital strategy should be fused with business strategy and owned by the "
    "management team, not delegated as a technical project (Bharadwaj et al., 2013; Westerman et "
    "al., 2014).",
    first_indent=0.30)
para(
    "For policy in Asian emerging economies, the same logic reframes public support. Because the "
    "binding constraint is capability rather than connectivity, support should shift from "
    "subsidising hardware towards building skills and absorptive capacity: management and "
    "digital-skills training, advisory and diagnostic services, and demonstrator projects that "
    "reduce the uncertainty about returns that deters SME investment. Two Asia-specific priorities "
    "follow. First, the finance constraint identified by Yoshino and Taghizadeh-Hesary (2016) should "
    "be addressed directly, by bundling affordable finance with advice and by building the SME "
    "databases and credit information that lower the cost of lending. Second, shared public digital "
    "infrastructure \u2014 interoperable digital-payment systems, digital identity and the "
    "e-government services on which Uzbek SME uptake is reportedly only about ten per cent (UNESCAP, "
    "2025) \u2014 lowers the fixed cost of transformation for the smallest firms and should be "
    "prioritised (OECD & ERIA, 2024; World Bank, 2023; Asian Development Bank, 2024).",
    first_indent=0.30)
para(
    "The two levels are complementary. Public investment in skills, finance and digital "
    "infrastructure raises the average capability level across the SME population, which increases "
    "the performance return to any given level of technology adoption and so strengthens the "
    "incentive for firms to invest; firm-level reinvestment of the resulting gains sustains the "
    "virtuous cycle. Where this complementarity is neglected \u2014 where infrastructure and grants "
    "are provided without capability \u2014 the predictable outcome is adoption without "
    "transformation, and the gap between dynamic and lagging firms persists.",
    first_indent=0.30)

# ===================== CONCLUSIONS =====================
heading("Conclusions")
para(
    "Across emerging Asia, the firms that most need the productivity gains of digital technology "
    "\u2014 the small and medium-sized enterprises that dominate employment and output \u2014 are "
    "also those least equipped to realise them. This study argued that the decisive factor is not "
    "access to technology but the managerial capability to absorb it. Integrating the process view "
    "of digital transformation with the dynamic-capabilities perspective, it set out a framework that "
    "links external drivers, the firm\u2019s sensing, seizing and reconfiguring capabilities, the "
    "depth of the transformation undertaken and the resulting performance, with a reinvestment "
    "feedback loop; it operationalised both digital-transformation intensity and dynamic capability "
    "using existing, validated instruments; and it expressed the core claim as a moderation "
    "hypothesis that can be estimated on firm-level data.",
    first_indent=0.30)
para(
    "Assessed in light of secondary evidence from the ASEAN region and Uzbekistan, the framework is "
    "consistent with a clear pattern: connectivity and policy ambition advance quickly, but the "
    "capability to turn technology into performance \u2014 skills, finance and management \u2014 lags "
    "behind. The practical message for managers is to sequence investment by capability and to own "
    "digital strategy at the top of the firm; for policymakers, to fund skills, finance and shared "
    "digital infrastructure rather than hardware alone. The study\u2019s main limitation is that the "
    "moderation is specified and illustrated rather than estimated; the numerical example uses "
    "didactic inputs, not data. The clear next step is to estimate equation (3) on firm-level "
    "evidence \u2014 the World Bank Enterprise Surveys or a primary survey of Uzbek and ASEAN SMEs "
    "\u2014 which would quantify the interaction \u03b2\u2083 and the capability threshold and sharpen "
    "the policy priorities.",
    first_indent=0.30)

# ===================== REFERENCES =====================
heading("References")
refs = [
    "Asian Development Bank. (2023). Asia small and medium-sized enterprise monitor 2023. Manila: "
    "Asian Development Bank. Available online: https://www.adb.org/publications/asia-sme-monitor-2023",

    "Asian Development Bank. (2024). Digital transformation for inclusive and sustainable "
    "development in Asia. Manila: Asian Development Bank. Available online: "
    "https://www.adb.org/publications/digital-transformation-for-inclusive-and-sustainable-development-in-asia",

    "Bharadwaj, A., El Sawy, O. A., Pavlou, P. A., & Venkatraman, N. (2013). Digital business "
    "strategy: Toward a next generation of insights. MIS Quarterly, 37(2), 471\u2013482. "
    "https://doi.org/10.25300/MISQ/2013/37:2.3",

    "Davis, F. D. (1989). Perceived usefulness, perceived ease of use, and user acceptance of "
    "information technology. MIS Quarterly, 13(3), 319\u2013340. https://doi.org/10.2307/249008",

    "de Mattos, C. S., Pellegrini, G., Hagelaar, G., & Dolfsma, W. (2024). Systematic literature "
    "review on technological transformation in SMEs: A transformation encompassing technology "
    "assimilation and business model innovation. Management Review Quarterly, 74(2), 1057\u20131095. "
    "https://doi.org/10.1007/s11301-023-00327-7",

    "Fitzgerald, M., Kruschwitz, N., Bonnet, D., & Welch, M. (2014). Embracing digital technology: A "
    "new strategic imperative. MIT Sloan Management Review, 55(2), 1\u201312.",

    "Hanelt, A., Bohnsack, R., Marz, D., & Antunes Marante, C. (2021). A systematic review of the "
    "literature on digital transformation: Insights and implications for strategy and organizational "
    "change. Journal of Management Studies, 58(5), 1159\u20131197. https://doi.org/10.1111/joms.12639",

    "Kraus, S., Durst, S., Ferreira, J. J., Veiga, P., Kailer, N., & Weinmann, A. (2022). Digital "
    "transformation in business and management research: An overview of the current status quo. "
    "International Journal of Information Management, 63, 102466. "
    "https://doi.org/10.1016/j.ijinfomgt.2021.102466",

    "Kump, B., Engelmann, A., Kessler, A., & Schweiger, C. (2019). Toward a dynamic capabilities "
    "scale: Measuring organizational sensing, seizing, and transforming capacities. Industrial and "
    "Corporate Change, 28(5), 1149\u20131172. https://doi.org/10.1093/icc/dty054",

    "Nambisan, S., Wright, M., & Feldman, M. (2019). The digital transformation of innovation and "
    "entrepreneurship: Progress, challenges and key themes. Research Policy, 48(8), 103773. "
    "https://doi.org/10.1016/j.respol.2019.03.018",

    "OECD. (2021). The digital transformation of SMEs. Paris: OECD Publishing. "
    "https://doi.org/10.1787/bdb9256a-en",

    "OECD, & ERIA. (2024). SME Policy Index: ASEAN 2024 \u2013 Enabling sustainable growth and "
    "digitalisation. Paris: OECD Publishing / Jakarta: Economic Research Institute for ASEAN and "
    "East Asia. https://doi.org/10.1787/f1f0c5f3-en",

    "Sagala, G. H., & \u0150ri, D. (2024). Toward SMEs digital transformation success: A systematic "
    "literature review. Information Systems and e-Business Management, 22(4), 667\u2013719. "
    "https://doi.org/10.1007/s10257-024-00682-2",

    "Teece, D. J. (2007). Explicating dynamic capabilities: The nature and microfoundations of "
    "(sustainable) enterprise performance. Strategic Management Journal, 28(13), 1319\u20131350. "
    "https://doi.org/10.1002/smj.640",

    "Teng, X., Wu, Z., & Yang, F. (2022). Research on the relationship between digital transformation "
    "and performance of SMEs. Sustainability, 14(10), 6012. https://doi.org/10.3390/su14106012",

    "Tornatzky, L. G., & Fleischer, M. (1990). The processes of technological innovation. Lexington, "
    "MA: Lexington Books.",

    "UNESCAP. (2025). Uzbekistan foresight on digital public services for small and medium-sized "
    "enterprises. Bangkok: United Nations Economic and Social Commission for Asia and the Pacific. "
    "Available online: https://www.unescap.org/events/2025/uzbekistan-foresight-digital-public-services-small-and-medium-sized-enterprises",

    "Venkatesh, V., Morris, M. G., Davis, G. B., & Davis, F. D. (2003). User acceptance of "
    "information technology: Toward a unified view. MIS Quarterly, 27(3), 425\u2013478. "
    "https://doi.org/10.2307/30036540",

    "Verhoef, P. C., Broekhuizen, T., Bart, Y., Bhattacharya, A., Qi Dong, J., Fabian, N., & "
    "Haenlein, M. (2021). Digital transformation: A multidisciplinary reflection and research "
    "agenda. Journal of Business Research, 122, 889\u2013901. "
    "https://doi.org/10.1016/j.jbusres.2019.09.022",

    "Vial, G. (2019). Understanding digital transformation: A review and a research agenda. The "
    "Journal of Strategic Information Systems, 28(2), 118\u2013144. "
    "https://doi.org/10.1016/j.jsis.2019.01.003",

    "Warner, K. S. R., & W\u00e4ger, M. (2019). Building dynamic capabilities for digital "
    "transformation: An ongoing process of strategic renewal. Long Range Planning, 52(3), "
    "326\u2013349. https://doi.org/10.1016/j.lrp.2018.12.001",

    "Westerman, G., Bonnet, D., & McAfee, A. (2014). Leading digital: Turning technology into "
    "business transformation. Boston, MA: Harvard Business Review Press.",

    "World Bank. (2023, November 30). World Bank to support Uzbekistan in developing the digital "
    "economy and creating new jobs in the information technology sector (Press release). Washington, "
    "DC: World Bank. Available online: "
    "https://www.worldbank.org/en/news/press-release/2023/11/30/world-bank-to-support-uzbekistan-in-developing-the-digital-economy-and-creating-new-jobs-in-the-it-sector",

    "World Bank. (2025, December 15). Improved access to finance to help 7,000 businesses in "
    "Uzbekistan grow and create jobs (Press release). Washington, DC: World Bank. Available online: "
    "https://www.worldbank.org/en/news/press-release/2025/12/15/improved-access-to-finance-to-help-7000-businesses-in-uzbekistan-grow-and-create-jobs",

    "Yoshino, N., & Taghizadeh-Hesary, F. (2016). Major challenges facing small and medium-sized "
    "enterprises in Asia and solutions for mitigating them (ADBI Working Paper No. 564). Tokyo: "
    "Asian Development Bank Institute. Available online: "
    "https://www.adb.org/publications/major-challenges-facing-small-and-medium-sized-enterprises-asia-and-solutions",
]
for ref in refs:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(4)
    pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.5)
    set_run(p.add_run(ref), size=14)

out = "Economic_Archive_Digital_Transformation_SMEs_Asia.docx"
doc.save(out)
print("Saved", out)
